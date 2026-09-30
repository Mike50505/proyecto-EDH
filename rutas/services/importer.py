"""Read legacy workbooks without running macros or trusting printed output sheets."""
import hashlib
import io
import tempfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation

import msoffcrypto
from django.db import transaction
from openpyxl import load_workbook

from rutas.models import BOMItem, Client, ImportIssue, ImportRun, Operation, Part, Route, SourceBook
from rutas.services.source_catalog import validate_source


class ImportErrorDetailed(Exception):
    pass


def clean(value):
    return "" if value is None else str(value).strip().replace("\xa0", " ").strip()


def raw(value):
    return "" if value is None else str(value)


def classify(filename):
    name = filename.lower()
    if "lennox" in name:
        return "General"
    if "slp" in name and "header" in name:
        return "SLP Headers"
    if "slp" in name and "individual" in name:
        return "SLP Individuales"
    if "header" in name:
        return "Headers"
    if "individual" in name:
        return "Individuales"
    raise ImportErrorDetailed("El nombre no identifica una variante Headers o Individuales; revise el archivo.")


def workbook_from_upload(upload, password=None):
    digest = hashlib.sha256()
    source = tempfile.SpooledTemporaryFile(max_size=16 * 1024 * 1024)
    for chunk in upload.chunks() if hasattr(upload, "chunks") else iter(lambda: upload.read(1024 * 1024), b""):
        source.write(chunk)
        digest.update(chunk)
    source.seek(0)
    signature = source.read(8)
    source.seek(0)
    if signature == bytes.fromhex("d0cf11e0a1b11ae1"):
        if not password:
            raise ImportErrorDetailed("El libro está cifrado. Introduzca la contraseña de apertura; no se guardará.")
        decrypted = tempfile.SpooledTemporaryFile(max_size=16 * 1024 * 1024)
        try:
            office = msoffcrypto.OfficeFile(source)
            office.load_key(password=password)
            office.decrypt(decrypted)
        except Exception as exc:
            raise ImportErrorDetailed("No se pudo abrir el libro cifrado. Verifique la contraseña y el archivo.") from exc
        decrypted.seek(0)
        source.close()
        source = decrypted
    elif signature[:2] != b"PK":
        raise ImportErrorDetailed("No es un archivo XLSM válido.")
    try:
        values = load_workbook(source, read_only=True, data_only=True, keep_vba=False, keep_links=False)
        source.seek(0)
        formulas = load_workbook(source, read_only=True, data_only=False, keep_vba=False, keep_links=False)
    except Exception as exc:
        source.close()
        raise ImportErrorDetailed("No se pudo leer la estructura XLSM.") from exc
    return digest.hexdigest(), source, values, formulas


def scan(upload, password=None, client_name=None, classification=None):
    digest, stream, wb, fw = workbook_from_upload(upload, password)
    try:
        required = {"Familias", "Ruta", "Datos"}
        if not required.issubset(set(wb.sheetnames)):
            raise ImportErrorDetailed("Faltan hojas requeridas: " + ", ".join(sorted(required - set(wb.sheetnames))))
        classification = classification or classify(upload.name)
        source_client = clean(wb["Datos"]["G2"].value)
        client_name = client_name or source_client or "DAIKIN"
        try:
            validate_source(client_name, classification)
        except ValueError as exc:
            raise ImportErrorDetailed(str(exc)) from exc
        families = []
        current_parent = None
        value_rows = wb["Familias"].iter_rows(min_row=2, max_row=5000, min_col=1, max_col=12, values_only=True)
        formula_rows = fw["Familias"].iter_rows(min_row=2, max_row=5000, min_col=1, max_col=12)
        for n, (row, formula_row) in enumerate(zip(value_rows, formula_rows), 2):
            code = clean(row[3])
            if not code:
                continue
            item = {"row": n, "type": clean(row[0]), "bom_revision": clean(row[1]), "drawing_revision": clean(row[2]), "code": code,
                    "description": raw(row[4]), "drawing": raw(row[5]), "quantity_raw": raw(row[6]), "od": raw(row[7]),
                    "wall": raw(row[8]), "development": raw(row[9]), "comments": raw(row[10]), "phase": raw(row[11]),
                    "formula": raw(formula_row[9].value) if formula_row[9].data_type == "f" else "",
                    "source_values": [raw(v) for v in row]}
            if classification in ("Headers", "SLP Headers", "General") and item["type"] == "FP":
                current_parent = n
            item["parent_row"] = current_parent if classification in ("Headers", "SLP Headers", "General") and current_parent != n else None
            families.append(item)
        header = next(fw["Ruta"].iter_rows(min_row=1, max_row=1, min_col=1, max_col=133, values_only=True))
        names = [clean(header[col - 1]) for col in range(2, 134, 4)]
        while names and not names[-1]:
            names.pop()
        routes = []
        for n, row in enumerate(wb["Ruta"].iter_rows(min_row=2, max_row=5000, min_col=1, max_col=1 + 4 * len(names), values_only=True), 2):
            original_code = raw(row[0])
            code = clean(original_code)
            if not code:
                continue
            operations = []
            for group, name in enumerate(names):
                offset = 1 + group * 4
                values = row[offset:offset + 4]
                if not any(clean(v) for v in values):
                    continue
                operations.append({"name": name or f"Grupo {group + 1} sin encabezado", "sequence": raw(values[0]),
                                   "tooling": raw(values[1]), "inspection": raw(values[2]), "machine": raw(values[3]), "group": group + 1})
            operations.sort(key=lambda op: (int(clean(op["sequence"])) if clean(op["sequence"]).isdigit() else 999999, op["group"]))
            routes.append({"row": n, "code": code, "original_code": original_code, "operations": operations})
        return {"sha256": digest, "classification": classification, "client": client_name, "source_client": source_client,
                "families": families, "routes": routes,
                "sheets": wb.sheetnames, "process_groups": names}
    finally:
        wb.close()
        fw.close()
        stream.close()


def import_workbook(upload, password=None, dry_run=True, user=None, client_name=None, classification=None):
    if not upload.name.lower().endswith(".xlsm"):
        raise ImportErrorDetailed("Solo se aceptan archivos .xlsm.")
    if getattr(upload, "size", 0) > 30 * 1024 * 1024:
        raise ImportErrorDetailed("El archivo supera el límite de 30 MB.")
    data = scan(upload, password, client_name, classification)
    conflicts = []
    if data["source_client"] and data["source_client"].casefold() != data["client"].casefold():
        conflicts.append(("Datos", 2, "", "cliente_distinto",
                          f"Datos!G2 dice {data['source_client']}; se usó el cliente seleccionado {data['client']}."))
    by_code = defaultdict(list)
    family_counts = Counter(item["code"] for item in data["families"])
    for route in data["routes"]:
        by_code[route["code"]].append(route)
        if family_counts[route["code"]] == 0:
            conflicts.append(("Ruta", route["row"], route["code"], "pieza_sin_familia",
                              "No existe una pieza con este código en Familias del mismo libro."))
        elif family_counts[route["code"]] > 1:
            conflicts.append(("Ruta", route["row"], route["code"], "pieza_ambigua",
                              "Hay varias piezas con este código en Familias del mismo libro."))
        sequences = [clean(op["sequence"]) for op in route["operations"] if clean(op["sequence"])]
        for seq, count in Counter(sequences).items():
            if count > 1:
                conflicts.append(("Ruta", route["row"], route["code"], "secuencia_repetida", f"La secuencia {seq} aparece {count} veces."))
        if not route["operations"]:
            conflicts.append(("Ruta", route["row"], route["code"], "sin_operaciones", "La fila no tiene operaciones; puede ser un ensamble."))
    for code, rows in by_code.items():
        if len(rows) > 1:
            for route in rows:
                conflicts.append(("Ruta", route["row"], code, "codigo_repetido", "Hay varias filas del mismo código en este libro; se conservan por separado."))
    for part in data["families"]:
        if not part["type"]:
            conflicts.append(("Familias", part["row"], part["code"], "tipo_vacio", "Clasifique el tipo de pieza."))
        if part["formula"] and not part["development"]:
            conflicts.append(("Familias", part["row"], part["code"], "formula_sin_valor", "La fórmula no tiene resultado almacenado."))
    summary = {"families": len(data["families"]), "route_rows": len(data["routes"]),
               "operations": sum(len(x["operations"]) for x in data["routes"]), "issues": len(conflicts),
               "classification": data["classification"], "client": data["client"], "sha256": data["sha256"], "dry_run": dry_run}
    with transaction.atomic():
        existing = SourceBook.objects.filter(sha256=data["sha256"]).first()
        if existing and (existing.classification != data["classification"] or
                         (existing.client_id and existing.client.name != data["client"])):
            raise ImportErrorDetailed("Este libro ya fue importado con otro cliente u origen.")
        if existing and not dry_run:
            summary["already_imported"] = True
        run = ImportRun.objects.create(source=existing, filename=upload.name[:255], sha256=data["sha256"], dry_run=dry_run, created_by=user, summary=summary)
        if dry_run or not existing:
            ImportIssue.objects.bulk_create([ImportIssue(run=run, sheet=s, row=r, code=c, kind=k, detail=d) for s, r, c, k, d in conflicts])
        if dry_run or existing:
            return run
        client, _ = Client.objects.get_or_create(name=data["client"])
        source = SourceBook.objects.create(client=client, filename=upload.name[:255], sha256=data["sha256"], classification=data["classification"])
        run.source = source
        run.save(update_fields=["source"])
        parts_by_row = {}
        parts_by_code = defaultdict(list)
        for item in data["families"]:
            part = Part.objects.create(client=client, code=item["code"], description=item["description"], part_type=item["type"],
                bom_revision=item["bom_revision"], drawing_revision=item["drawing_revision"], drawing_number=item["drawing"],
                od_raw=item["od"], wall_raw=item["wall"], development_raw=item["development"], development_formula=item["formula"],
                comments=item["comments"], phase=item["phase"], source=source, source_row=item["row"],
                source_values=item["source_values"], needs_review=not item["type"])
            parts_by_row[item["row"]] = part
            parts_by_code[item["code"]].append(part)
        for item in data["families"]:
            if item["parent_row"] and item["parent_row"] in parts_by_row:
                try:
                    quantity = Decimal(clean(item["quantity_raw"]) or "0")
                except InvalidOperation:
                    quantity = Decimal("0")
                    ImportIssue.objects.create(run=run, sheet="Familias", row=item["row"], code=item["code"], kind="cantidad_ambigua", detail="La cantidad BOM no pudo interpretarse; revise antes de emitir.")
                BOMItem.objects.create(parent=parts_by_row[item["parent_row"]], component=parts_by_row[item["row"]],
                                       position=item["row"], quantity_per=quantity, source_raw_quantity=item["quantity_raw"])
        conflicted = {(r, c) for s, r, c, k, d in conflicts if s == "Ruta"}
        for item in data["routes"]:
            matching = parts_by_code[item["code"]]
            review = not source.print_approved or (item["row"], item["code"]) in conflicted or len(matching) != 1
            route = Route.objects.create(client=client, part=matching[0] if len(matching) == 1 else None,
                code=item["code"], description=matching[0].description if len(matching) == 1 else "",
                classification=data["classification"], status=Route.REVIEW if review else Route.ACTIVE,
                source=source, source_sheet="Ruta", source_row=item["row"], source_code=item["original_code"],
                source_values={"original_code": item["original_code"]})
            Operation.objects.bulk_create([Operation(route=route, position=i, source_sequence=op["sequence"], name=op["name"],
                tooling=op["tooling"], inspection=op["inspection"], machine=op["machine"], source_group=op["group"])
                for i, op in enumerate(item["operations"], 1)])
        summary["imported_route_rows"] = len(data["routes"])
        run.summary = summary
        run.save(update_fields=["summary"])
        return run
