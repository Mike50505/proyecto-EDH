"""Transient Excel-like order selection. Uploaded rows never reach the database."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook
from django.utils import timezone

from rutas.models import Route
from rutas.services.catalog import parent_routes
from rutas.services.documents import preview_labels_for_line, render_pdf, validate_label_mode

HEADERS = ("ITEM PADRE", "CANTIDAD", "SEMANA")
PREVIOUS_HEADERS = ("ITEM PADRE", "CANTIDAD", "SEMANA", "RE")
LEGACY_HEADERS = ("SHOP ORDER", "ITEM PADRE", "CANTIDAD", "SEMANA", "RE", "SECUENCIA")
OLD_COMMON_HEADERS = ("LINEA", "PLANNER", "RESPONSABLE")
FIELDS = ("shop_order", "parent_code", "quantity", "week", "re", "sequence", "line", "planner", "responsible", "route_id")
MAX_ROWS = 200


def value_text(value):
    return "" if value is None else str(value).strip()


def blank_rows(count=12):
    return [{field: "" for field in FIELDS} for _ in range(count)]


def read_excel(upload):
    if not upload.name.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Carga un archivo Excel .xlsx o .xlsm.")
    if upload.size > 2 * 1024 * 1024:
        raise ValueError("El archivo supera 2 MB. Usa la plantilla solo para seleccionar hasta 200 órdenes.")
    try:
        book = load_workbook(BytesIO(upload.read()), read_only=True, data_only=True)
        sheet = book.active
        headings = tuple(value_text(cell.value).upper() for cell in sheet[1][:9])
        legacy = headings[:6] == LEGACY_HEADERS and (not any(headings[6:]) or headings[6:] == OLD_COMMON_HEADERS)
        previous = headings[:4] == PREVIOUS_HEADERS and not any(headings[4:])
        current = headings[:3] == HEADERS and not any(headings[3:])
        if not current and not previous and not legacy:
            raise ValueError("La fila 1 debe contener ITEM PADRE, CANTIDAD y SEMANA, en ese orden.")
        rows = []
        for cells in sheet.iter_rows(min_row=2, max_col=9 if legacy else 4 if previous else 3, values_only=True):
            if legacy:
                row = {field: value_text(value) for field, value in zip(FIELDS, cells)}
            else:
                row = {field: "" for field in FIELDS}
                for field, value in zip(("parent_code", "quantity", "week", "re") if previous else
                                        ("parent_code", "quantity", "week"), cells):
                    row[field] = value_text(value)
            if any(row[field] for field in ("parent_code", "quantity", "week")):
                row["re"] = ""
                row["shop_order"] = row["week"]
                row["sequence"] = str(len(rows) + 1)
                rows.append(row)
            if len(rows) > MAX_ROWS:
                raise ValueError("La plantilla admite hasta 200 órdenes por lote.")
        if not rows:
            raise ValueError("La plantilla no contiene órdenes.")
        return rows
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("No se pudo leer el Excel. Descarga la plantilla y verifica que el archivo no esté dañado o cifrado.") from exc


def rows_from_post(post):
    try:
        count = int(post.get("row_count", 0))
    except (TypeError, ValueError) as exc:
        raise ValueError("Número de filas inválido.") from exc
    if count < 1 or count > MAX_ROWS:
        raise ValueError("El lote debe tener entre 1 y 200 filas.")
    rows = []
    for index in range(count):
        row = {field: value_text(post.get(f"rows-{index}-{field}")) for field in FIELDS}
        if any(row.get(field) for field in ("parent_code", "quantity", "week")):
            row["shop_order"] = row["week"]
            row["sequence"] = str(len(rows) + 1)
            rows.append(row)
    if not rows:
        raise ValueError("Agrega al menos una orden.")
    for field in ("line", "planner", "responsible"):
        rows[0][field] = value_text(post.get(f"common_{field}"))
    return rows


def build_pdf(rows, template_key="mesa-two-labels-v1", label_mode="components_only"):
    """Resolve each row against the whole catalog and render entirely in memory."""
    validate_label_mode(label_mode)
    first = rows[0]
    defaults = {key: first.get(key, "") for key in ("week", "line", "planner", "responsible")}
    labels = []
    errors = []
    clients = set()
    has_unapproved_routes = False
    for row_index, row in enumerate(rows, 1):
        display_row = row_index + 1
        effective = {"week": row.get("week") or defaults["week"],
                     "line": defaults["line"], "planner": defaults["planner"],
                     "responsible": defaults["responsible"]}
        if not row.get("parent_code"):
            errors.append(f"Fila {display_row}: falta ITEM PADRE.")
        if not effective["week"]:
            errors.append(f"Fila {display_row}: falta SEMANA.")
        try:
            quantity = Decimal(row["quantity"])
            if not quantity.is_finite() or quantity <= 0:
                raise InvalidOperation
        except (InvalidOperation, KeyError):
            errors.append(f"Fila {display_row}: CANTIDAD debe ser mayor que cero.")
            continue
        sequence = row_index
        if not row.get("parent_code"):
            continue
        routes = parent_routes().filter(code__iexact=row["parent_code"]).exclude(
            status=Route.ARCHIVED).select_related("client", "part", "source")
        matches = list(routes.order_by("client__name", "classification", "source_row", "id")[:2])
        if not matches:
            errors.append(f"Fila {display_row}: no se encontró ITEM PADRE {row['parent_code']} en el catálogo.")
            continue
        selected_id = str(row.get("route_id") or "").strip()
        if selected_id:
            valid_id = selected_id.isascii() and selected_id.isdigit() and len(selected_id) <= 18
            route = routes.filter(pk=selected_id).first() if valid_id else None
            if route is None:
                errors.append(f"Fila {display_row}: la ruta elegida no corresponde a {row['parent_code']}.")
                continue
        elif len(matches) == 1:
            route = matches[0]
        else:
            errors.append(f"Fila {display_row}: {row['parent_code']} tiene varias rutas. Elige cliente y origen en esa fila.")
            continue
        clients.add(route.client.name)
        line = SimpleNamespace(route_id=route.pk, shop_order=effective["week"], quantity=quantity, position=sequence)
        try:
            route_labels, _, needs_review = preview_labels_for_line(line, label_mode)
        except ValueError as exc:
            errors.append(f"Fila {display_row}: {exc}")
            continue
        has_unapproved_routes = has_unapproved_routes or needs_review
        for label in route_labels:
            label.update(effective)
            label["client"] = route.client.name
        labels.extend(route_labels)
        row["re"] = str(route_labels[0]["re_count"])
    if errors:
        raise ValueError("\n".join(errors))
    snapshot = {
        "client": next(iter(clients)) if len(clients) == 1 else "Varios clientes", "week": defaults["week"], "line": defaults["line"],
        "planner": defaults["planner"], "responsible": defaults["responsible"],
        "issue_date": timezone.localdate().isoformat(),
        "ship_date": "", "copies": 1, "lines": [], "labels": labels,
        "label_mode": label_mode,
        "preview_status": "VISTA PREVIA · NO APROBADA" if has_unapproved_routes else "SELECCIÓN TEMPORAL · NO GUARDADA",
    }
    return render_pdf(snapshot, template_key)
