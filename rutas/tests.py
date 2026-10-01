import io
import tempfile
from decimal import Decimal

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

from rutas.forms import ScheduleLineForm
from rutas.models import BOMItem, Client, ImportIssue, ImportRun, IssuedDocument, Operation, Part, Route, Schedule, ScheduleLine, SourceBook
from rutas.services.documents import issue_schedule, render_pdf
from rutas.services.importer import ImportErrorDetailed, import_workbook
from rutas.services.reconciliation import reconcile_all


def example_xlsm():
    wb = Workbook()
    family = wb.active
    family.title = "Familias"
    family.append(["Tipo", "BOM", "Dibujo rev", "Parte", "Descripción", "Número dibujo", "Cantidad", "OD", "Pared", "Desarrollo", "Comentario", "Fase"])
    family.append(["FP", "A", "B", "0210A00205", "Pieza de prueba", "PL-1", 1, "3/8", "7mm", "223*", "", ""])
    route = wb.create_sheet("Ruta")
    route.cell(1, 2, "CORTE")
    route.cell(1, 6, "DOBLEZ 1")
    route.cell(2, 1, "0210A00205")
    route.cell(2, 2, 2)
    route.cell(2, 3, "H-1")
    route.cell(2, 4, "I-1")
    route.cell(2, 5, "M-1")
    route.cell(2, 6, 1)
    route.cell(3, 1, "\xa00210A00205 ")
    route.cell(3, 2, 1)
    wb.create_sheet("Datos").cell(2, 7, "DAIKIN")
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile("Macro Headers.xlsm", out.getvalue())


class ImportTests(TestCase):
    def test_encrypted_file_requires_password_without_leaking_it(self):
        upload = SimpleUploadedFile("Macro Headers.xlsm", bytes.fromhex("d0cf11e0a1b11ae1") + b"test")
        with self.assertRaisesRegex(ImportErrorDetailed, "cifrado"):
            import_workbook(upload)

    def test_dry_run_and_repeat_keep_duplicate_rows_for_review(self):
        run = import_workbook(example_xlsm(), dry_run=True)
        self.assertEqual(run.summary["route_rows"], 2)
        self.assertEqual(Route.objects.count(), 0)
        self.assertEqual(SourceBook.objects.count(), 0)
        self.assertTrue(ImportIssue.objects.filter(kind="codigo_repetido").exists())
        import_workbook(example_xlsm(), dry_run=False)
        again = import_workbook(example_xlsm(), dry_run=False)
        self.assertTrue(again.summary["already_imported"])
        self.assertEqual(again.issues.count(), 0)
        self.assertEqual(Route.objects.count(), 2)
        self.assertEqual(list(Route.objects.values_list("source_row", flat=True).order_by("source_row")), [2, 3])
        self.assertEqual(Route.objects.filter(status=Route.REVIEW).count(), 2)
        self.assertEqual(Route.objects.first().code, "0210A00205")
        self.assertEqual(Part.objects.get(source_row=2).quantity_raw, "1")
        self.assertEqual(Route.objects.get(source_row=2).operations.first().name, "DOBLEZ 1")

    def test_selected_client_overrides_misspelled_workbook_cell_and_stays_in_review(self):
        workbook = example_xlsm()
        run = import_workbook(workbook, password=None, dry_run=False, client_name="RHEEM", classification="Headers")
        self.assertEqual(run.source.client.name, "RHEEM")
        self.assertFalse(run.source.print_approved)
        self.assertEqual(Route.objects.filter(client__name="RHEEM", status=Route.REVIEW).count(), 2)
        self.assertTrue(run.issues.filter(kind="cliente_distinto").exists())

    def test_import_rejects_invalid_client_variant_pair(self):
        with self.assertRaisesRegex(ImportErrorDetailed, "combinación válida"):
            import_workbook(example_xlsm(), dry_run=True, client_name="LENNOX", classification="Headers")

    def test_issue_list_excludes_dry_run_duplicates(self):
        simulation = import_workbook(example_xlsm(), dry_run=True)
        committed = import_workbook(example_xlsm(), dry_run=False)
        reviewer = get_user_model().objects.create_superuser("revisor", "r@example.test", "example-long-password")
        self.client.force_login(reviewer)
        response = self.client.get(reverse("rutas:issues"))
        self.assertEqual(response.status_code, 200)
        visible_ids = {issue.pk for issue in response.context["page"]}
        self.assertTrue(visible_ids)
        self.assertTrue(visible_ids.issubset(set(committed.issues.values_list("pk", flat=True))))
        self.assertFalse(visible_ids.intersection(simulation.issues.values_list("pk", flat=True)))


class RouteEditorTests(TestCase):
    def test_cell_editor_and_form_editor_share_validated_operation_fields(self):
        user = get_user_model().objects.create_superuser("capturista", "c@example.test", "example-long-password")
        client = Client.objects.create(name="RHEEM")
        self.client.force_login(user)
        page = self.client.get(reverse("rutas:create"))
        self.assertContains(page, 'data-editor-switch="sheet"')
        self.assertContains(page, 'data-editor-switch="form"')
        self.assertContains(page, 'data-family-editor')
        self.assertContains(page, 'data-family-paste')
        self.assertEqual(page.content.count(b'name="family-code"'), 1)
        self.assertContains(page, "Número dibujo")
        prefix = page.context["formset"].prefix
        data = {"client": client.pk, "family-part_type": "FP", "family-code": "R-PRUEBA",
                "family-description": "Prueba", "family-bom_revision": "A", "family-drawing_revision": "B",
                "family-drawing_number": "PL-1", "family-quantity_raw": "3", "family-od_raw": "3/8",
                "family-wall_raw": "0.028", "family-development_raw": "223*",
                "family-comments": "Revisar", "family-phase": "CORTE",
                "classification": "Headers", "revision": "", "status": Route.REVIEW, "editor_mode": "sheet",
                f"{prefix}-TOTAL_FORMS": "2", f"{prefix}-INITIAL_FORMS": "0",
                f"{prefix}-MIN_NUM_FORMS": "0", f"{prefix}-MAX_NUM_FORMS": "200"}
        for index, name in enumerate(("CORTE", "DOBLEZ")):
            data.update({f"{prefix}-{index}-position": str(index + 1),
                         f"{prefix}-{index}-source_sequence": str(index + 1),
                         f"{prefix}-{index}-name": name,
                         f"{prefix}-{index}-machine": f"M-{index + 1}"})
        response = self.client.post(reverse("rutas:create"), data)
        self.assertEqual(response.status_code, 302)
        route = Route.objects.get(code="R-PRUEBA")
        self.assertEqual(route.description, "Prueba")
        self.assertEqual(route.part.client, client)
        self.assertEqual(route.part.quantity_raw, "3")
        self.assertEqual(route.part.development_raw, "223*")
        self.assertIsNone(route.part.source_id)
        self.assertEqual(list(route.operations.values_list("name", flat=True)), ["CORTE", "DOBLEZ"])
        self.assertEqual(route.operations.count(), 2)

        data["family-code"] = "R-INVALIDA"
        data[f"{prefix}-1-position"] = "1"
        data["editor_mode"] = "form"
        response = self.client.post(reverse("rutas:create"), data)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'data-editor-mode="form"')
        self.assertFalse(response.context["formset"].is_valid())
        self.assertFalse(Route.objects.filter(code="R-INVALIDA").exists())
        self.assertFalse(Part.objects.filter(code="R-INVALIDA").exists())

    def test_existing_operation_can_be_edited_in_cells(self):
        user = get_user_model().objects.create_superuser("editor", "e@example.test", "example-long-password")
        client = Client.objects.create(name="LENNOX")
        part = Part.objects.create(client=client, code="L-PIEZA")
        route = Route.objects.create(client=client, part=part, code="L-PRUEBA", status=Route.REVIEW)
        operation = Operation.objects.create(route=route, position=1, source_sequence="1", name="CORTE")
        self.client.force_login(user)
        url = reverse("rutas:edit", args=[route.pk])
        page = self.client.get(url)
        prefix = page.context["formset"].prefix
        self.assertContains(page, f'name="{prefix}-0-id" value="{operation.pk}"')
        self.assertContains(page, f'data-client-id="{client.pk}" data-part-code="L-PIEZA"')
        self.assertContains(page, "L-PIEZA · Manual")
        data = {"client": client.pk, "part": part.pk, "code": route.code, "description": "",
                "classification": "General", "revision": "", "status": Route.REVIEW, "version": route.version,
                "editor_mode": "sheet", f"{prefix}-TOTAL_FORMS": "1", f"{prefix}-INITIAL_FORMS": "1",
                f"{prefix}-MIN_NUM_FORMS": "0", f"{prefix}-MAX_NUM_FORMS": "200",
                f"{prefix}-0-id": operation.pk, f"{prefix}-0-position": "1",
                f"{prefix}-0-source_sequence": "1", f"{prefix}-0-name": "DOBLEZ",
                f"{prefix}-0-machine": "M-1"}
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 302)
        self.assertEqual(route.operations.count(), 1)
        self.assertEqual(route.operations.first().name, "DOBLEZ")


class PdfPaginationTests(TestCase):
    def test_long_operations_stay_two_labels_per_page(self):
        operations = [{"position": index, "sequence": str(index), "name": "PERFORACIÓN Y DOBLEZ",
                       "machine": "MÁQUINA DE PRODUCCIÓN", "tooling": "RODILLOS, MORDAZAS Y GUÍA DE AJUSTE 5/8",
                       "inspection": "LIBERACIÓN DE PRIMERA PIEZA"} for index in range(1, 11)]
        labels = [{"parent_code": f"P-{index}", "parent_description": "REFRIGERANT PIPE ASSY (GAS)",
                   "component_code": f"C-{index}", "component_description": "REFRIGERANT PIPE (GAS)",
                   "dimensions": {"od": "0.625", "wall": "0.037", "development": "515"},
                   "shop_order": f"RAMOS {index}", "quantity": "40", "sequence": index,
                   "label_index": 2, "label_total": 3, "re_count": 2, "classification": "Headers",
                   "operations": operations} for index in range(1, 13)]
        snapshot = {"client": "DAIKIN", "week": "31A", "line": "DAIKIN", "planner": "Juanjo",
                    "responsible": "Edith", "issue_date": "2026-09-30", "ship_date": "", "copies": 1,
                    "labels": labels, "preview_status": "NO APROBADA"}
        pages = PdfReader(io.BytesIO(render_pdf(snapshot))).pages
        self.assertEqual(len(pages), 6)
        self.assertTrue(all(page.extract_text().count("ETIQUETA DE CORTE") == 2 for page in pages))


class DocumentTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.temp.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = get_user_model().objects.create_user("lector", password="example-long-password")
        self.client_record = Client.objects.create(name="DAIKIN")
        self.route = Route.objects.create(client=self.client_record, code="0210A00205", description="Pieza inicial", status=Route.ACTIVE)
        Operation.objects.create(route=self.route, position=1, source_sequence="3", name="CORTE", tooling="H-1", inspection="I-1", machine="M-1")
        self.schedule = Schedule.objects.create(client=self.client_record, week="31A", planner="Ana")
        ScheduleLine.objects.create(schedule=self.schedule, route=self.route, shop_order="SO-A", quantity=Decimal("2"), position=1)

    def test_historical_pdf_snapshot_survives_route_edit(self):
        document = issue_schedule(self.schedule, self.user)
        self.assertTrue(document.pdf.open("rb").read(4) == b"%PDF")
        self.route.description = "Descripción modificada"
        self.route.save()
        document.refresh_from_db()
        self.assertEqual(document.snapshot["lines"][0]["description"], "Pieza inicial")
        self.assertEqual(document.snapshot["week"], "31A")
        self.assertTrue(document.snapshot["issue_date"])
        self.assertEqual(document.snapshot["lines"][0]["operations"][0]["inspection"], "I-1")

    def test_document_download_checks_server_permission(self):
        document = issue_schedule(self.schedule, self.user)
        self.client.force_login(self.user)
        url = reverse("rutas:document", args=[document.pk])
        self.assertEqual(self.client.get(url).status_code, 403)
        permission = Permission.objects.get(codename="view_issueddocument")
        self.user.user_permissions.add(permission)
        self.assertEqual(self.client.get(url).status_code, 200)

    def test_review_route_cannot_be_issued(self):
        self.route.status = Route.REVIEW
        self.route.save()
        with self.assertRaisesRegex(ValueError, "requiere revisión"):
            issue_schedule(self.schedule, self.user)

    def test_unapproved_source_cannot_be_issued(self):
        source = SourceBook.objects.create(client=self.client_record, filename="nuevo.xlsm", sha256="a" * 64,
                                           classification="Headers", print_approved=False)
        self.route.source = source
        self.route.save()
        with self.assertRaisesRegex(ValueError, "formato de impresión"):
            issue_schedule(self.schedule, self.user)

    def test_copies_repeat_document_without_changing_piece_quantity(self):
        self.schedule.copies = 2
        self.schedule.save()
        document = issue_schedule(self.schedule, self.user)
        self.assertEqual(document.snapshot["copies"], 2)
        self.assertEqual(Decimal(document.snapshot["lines"][0]["quantity"]), Decimal("2"))
        self.assertEqual(len(PdfReader(io.BytesIO(document.pdf.open("rb").read())).pages), 2)

    def test_parent_prints_two_rm_labels_on_one_sheet(self):
        parent_part = Part.objects.create(client=self.client_record, code=self.route.code, part_type="FP")
        self.route.part = parent_part
        self.route.save(update_fields=["part"])
        for position, (code, per_unit) in enumerate((("RM-A", "2"), ("RM-B", "3")), 1):
            component = Part.objects.create(client=self.client_record, code=code, part_type="RM")
            BOMItem.objects.create(parent=parent_part, component=component, position=position, quantity_per=Decimal(per_unit))
            child = Route.objects.create(client=self.client_record, part=component, code=code, status=Route.ACTIVE)
            Operation.objects.create(route=child, position=1, name="DOBLEZ", machine=f"M-{position}")
        unused = Part.objects.create(client=self.client_record, code="BR-A", part_type="BR")
        BOMItem.objects.create(parent=parent_part, component=unused, position=3, quantity_per=Decimal("4"))

        document = issue_schedule(self.schedule, self.user)
        labels = document.snapshot["labels"]
        self.assertEqual([label["component_code"] for label in labels], ["RM-A", "RM-B"])
        self.assertEqual([label["label_index"] for label in labels], [2, 3])
        self.assertEqual([label["re_count"] for label in labels], [2, 2])
        self.assertEqual([label["quantity"] for label in labels], ["4", "6"])
        self.assertEqual(len(PdfReader(io.BytesIO(document.pdf.open("rb").read())).pages), 1)
        self.assertEqual(ScheduleLineForm(initial={"route": self.route.pk}).fields["re_count"].initial, 2)

    def test_parent_requires_unique_active_component_route(self):
        parent_part = Part.objects.create(client=self.client_record, code=self.route.code, part_type="FP")
        component = Part.objects.create(client=self.client_record, code="RM-A", part_type="RM")
        BOMItem.objects.create(parent=parent_part, component=component, position=1, quantity_per=Decimal("1"))
        self.route.part = parent_part
        self.route.save(update_fields=["part"])
        with self.assertRaisesRegex(ValueError, "ruta activa única"):
            issue_schedule(self.schedule, self.user)

    def test_catalog_and_detail_render_for_signed_in_user(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("rutas:list")).status_code, 200)
        self.assertContains(self.client.get(reverse("rutas:detail", args=[self.route.pk])), "CORTE")

    def test_print_button_opens_route_preview_without_issuing_document(self):
        self.user.user_permissions.add(Permission.objects.get(codename="view_route"))
        self.client.force_login(self.user)
        print_url = reverse("rutas:route_print", args=[self.route.pk])
        self.assertContains(self.client.get(reverse("rutas:list")), print_url)
        self.assertContains(self.client.get(reverse("rutas:detail", args=[self.route.pk])), "Imprimir ruta")
        response = self.client.get(print_url)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertEqual(response["Cache-Control"], "private, no-store")
        pdf = b"".join(response.streaming_content)
        self.assertEqual(pdf[:4], b"%PDF")
        self.assertIn("NO PRODUCCIÓN", PdfReader(io.BytesIO(pdf)).pages[0].extract_text())
        self.assertEqual(IssuedDocument.objects.count(), 0)

    def test_print_preview_requires_view_permission(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse("rutas:route_print", args=[self.route.pk])).status_code, 403)

    def test_page_selection_prefills_schedule(self):
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"))
        self.client.force_login(self.user)
        response = self.client.post(reverse("rutas:select"), {"scope": "page", "route_ids": [self.route.pk]})
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Seleccionar órdenes")
        self.assertContains(response, 'value="0210A00205"')
        self.assertEqual(Schedule.objects.count(), 1)

    def test_batch_grid_suggests_parent_items(self):
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"))
        self.client.force_login(self.user)
        response = self.client.get(reverse("rutas:batch_print"))
        self.assertContains(response, '<datalist id="parent-code-options">')
        self.assertContains(response, '<option value="0210A00205"></option>')
        self.assertContains(response, 'list="parent-code-options" data-parent-code autocomplete="off"')
        self.assertNotContains(response, '<th>SHOP ORDER</th>')
        self.assertNotContains(response, '<th>Secuencia</th>')

    def test_client_filter_and_batch_pdf_use_selected_client(self):
        rheem = Client.objects.create(name="RHEEM")
        rheem_route = Route.objects.create(client=rheem, code=self.route.code, status=Route.ACTIVE,
                                           classification="Individuales", description="Ruta RHEEM")
        Operation.objects.create(route=rheem_route, position=1, name="PERFORACIÓN")
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"),
                                       Permission.objects.get(codename="add_issueddocument"))
        self.client.force_login(self.user)
        page = self.client.get(reverse("rutas:list"), {"client": "RHEEM"})
        self.assertContains(page, "Ruta RHEEM")
        self.assertNotContains(page, "Pieza inicial")
        grid = self.client.get(reverse("rutas:batch_print"), {"client": "RHEEM", "classification": "Individuales"})
        self.assertContains(grid, 'value="RHEEM" selected')
        response = self.client.post(reverse("rutas:batch_print"), {"action": "print", "client": "RHEEM",
            "classification": "Individuales", "row_count": "1", "rows-0-shop_order": "RH-1",
            "rows-0-parent_code": rheem_route.code, "rows-0-quantity": "1", "rows-0-week": "31A",
            "rows-0-sequence": "1", "common_line": "RHEEM", "common_planner": "Ana",
            "common_responsible": "Luis"})
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("RHEEM", PdfReader(io.BytesIO(b"".join(response.streaming_content))).pages[0].extract_text())

    def test_excel_upload_and_grid_print_do_not_save_orders(self):
        second = Route.objects.create(client=self.client_record, code="RM-B", status=Route.ACTIVE)
        Operation.objects.create(route=second, position=1, name="INSPECCIÓN")
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"),
                                       Permission.objects.get(codename="add_issueddocument"))
        self.client.force_login(self.user)
        template_response = self.client.get(reverse("rutas:batch_template"))
        self.assertEqual(template_response.status_code, 200)
        template_book = load_workbook(io.BytesIO(b"".join(template_response.streaming_content)), read_only=True)
        self.assertEqual([template_book.active.cell(1, col).value for col in range(1, 4)],
                         ["ITEM PADRE", "CANTIDAD", "SEMANA"])
        self.assertIsNone(template_book.active["D1"].value)
        wb = Workbook()
        sheet = wb.active
        sheet.append(["SHOP ORDER", "ITEM PADRE", "CANTIDAD", "SEMANA", "RE", "Secuencia", "LINEA", "PLANNER", "RESPONSABLE"])
        sheet.append(["SO-1", self.route.code, 10, "31A", 0, 1, "DAIKIN", "Ana", "Luis"])
        sheet.append(["SO-2", second.code, 20, "31A", 0, 2, "OTRA", "Otro", "Otra"])
        stream = io.BytesIO()
        wb.save(stream)
        upload = SimpleUploadedFile("ordenes.xlsx", stream.getvalue())
        url = reverse("rutas:batch_print")
        response = self.client.post(url, {"action": "upload", "workbook": upload})
        self.assertContains(response, 'value="0210A00205"')
        self.assertNotContains(response, 'name="rows-0-shop_order"')
        self.assertEqual(response.context["rows"][0]["shop_order"], "31A")
        self.assertEqual(response.context["rows"][1]["sequence"], "2")
        self.assertContains(response, 'name="common_line" value="DAIKIN"')
        self.assertContains(response, 'name="common_planner" value="Ana"')
        self.assertEqual(response.content.count(b'name="common_line"'), 1)
        self.assertNotIn(b'<th>LINEA</th>', response.content)
        self.assertNotIn(b'<th>PLANNER</th>', response.content)
        self.assertNotIn(b'<th>RESPONSABLE</th>', response.content)
        self.assertNotIn(b'name="rows-1-line"', response.content)
        self.assertEqual(Schedule.objects.count(), 1)
        self.assertEqual(IssuedDocument.objects.count(), 0)

        short_book = Workbook()
        short_book.active.append(["ITEM PADRE", "CANTIDAD", "SEMANA"])
        short_book.active.append([self.route.code, 5, "32B"])
        short_stream = io.BytesIO()
        short_book.save(short_stream)
        short_upload = SimpleUploadedFile("seleccion.xlsx", short_stream.getvalue())
        short_response = self.client.post(url, {"action": "upload", "workbook": short_upload,
                                                 "common_line": "DAIKIN", "common_planner": "Ana",
                                                 "common_responsible": "Luis"})
        self.assertEqual(short_response.context["rows"][0]["shop_order"], "32B")
        self.assertEqual(short_response.context["rows"][0]["sequence"], "1")
        self.assertContains(short_response, 'name="common_responsible" value="Luis"')

        data = {"action": "print", "row_count": "2", "classification": "",
                "rows-0-shop_order": "SO-1", "rows-0-parent_code": self.route.code,
                "rows-0-quantity": "10", "rows-0-week": "31A", "rows-0-re": "0",
                "rows-0-sequence": "1", "common_line": "DAIKIN", "common_planner": "Ana",
                "common_responsible": "Luis", "rows-1-shop_order": "SO-2",
                "rows-1-parent_code": second.code, "rows-1-quantity": "20",
                "rows-1-sequence": "2"}
        response = self.client.post(url, data)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        pdf = b"".join(response.streaming_content)
        self.assertEqual(pdf[:4], b"%PDF")
        self.assertEqual(len(PdfReader(io.BytesIO(pdf)).pages), 1)
        self.assertIn("RM-B", PdfReader(io.BytesIO(pdf)).pages[0].extract_text())
        self.assertNotIn("SO-1", PdfReader(io.BytesIO(pdf)).pages[0].extract_text())
        self.assertNotIn("SO-2", PdfReader(io.BytesIO(pdf)).pages[0].extract_text())
        self.assertEqual(PdfReader(io.BytesIO(pdf)).pages[0].extract_text().count("Ana"), 2)
        self.assertEqual(Schedule.objects.count(), 1)
        self.assertEqual(IssuedDocument.objects.count(), 0)

    def test_batch_print_allows_marked_preview_of_unapproved_route(self):
        self.route.status = Route.REVIEW
        self.route.save()
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"),
                                       Permission.objects.get(codename="add_issueddocument"))
        self.client.force_login(self.user)
        response = self.client.post(reverse("rutas:batch_print"), {"action": "print", "row_count": "1",
            "rows-0-shop_order": "SO-1", "rows-0-parent_code": self.route.code,
            "rows-0-quantity": "10", "rows-0-week": "31A", "rows-0-sequence": "1",
            "common_line": "DAIKIN", "common_planner": "Ana", "common_responsible": "Luis"})
        self.assertEqual(response["Content-Type"], "application/pdf")
        text = PdfReader(io.BytesIO(b"".join(response.streaming_content))).pages[0].extract_text()
        self.assertIn("NO APROBADA", text)
        self.assertEqual(IssuedDocument.objects.count(), 0)

    def test_batch_print_has_all_datos_fields_and_one_pdf(self):
        second = Route.objects.create(client=self.client_record, code="RM-B", status=Route.ACTIVE)
        Operation.objects.create(route=second, position=1, name="INSPECCIÓN")
        self.user.user_permissions.add(Permission.objects.get(codename="add_schedule"),
                                       Permission.objects.get(codename="add_issueddocument"),
                                       Permission.objects.get(codename="view_issueddocument"))
        self.client.force_login(self.user)
        session = self.client.session
        session["selected_route_ids"] = [self.route.pk, second.pk]
        session.save()
        page = self.client.get(reverse("rutas:schedule_create"))
        for label in ("CANTIDAD", "SEMANA", "RE (calculado)", "Secuencia", "LINEA", "PLANNER", "RESPONSABLE"):
            self.assertContains(page, label)
        self.assertContains(page, "Generar PDF de todas")
        self.assertContains(page, f'value="{second.pk}" selected')

        data = {"client": self.client_record.pk, "week": "31A", "line": "L1", "planner": "Ana",
                "responsible": "Luis", "issue_date": "2026-09-29", "ship_date": "", "copies": "1",
                "action": "print_all", "lines-TOTAL_FORMS": "2", "lines-INITIAL_FORMS": "0",
                "lines-MIN_NUM_FORMS": "0", "lines-MAX_NUM_FORMS": "1000"}
        for index, route in enumerate((self.route, second)):
            data.update({f"lines-{index}-position": str(index + 1), f"lines-{index}-route": str(route.pk),
                         f"lines-{index}-shop_order": f"SO-{index + 1}", f"lines-{index}-quantity": "2"})
        response = self.client.post(reverse("rutas:schedule_create"), data)
        self.assertEqual(response.status_code, 302)
        document = IssuedDocument.objects.get()
        self.assertEqual(response.url, reverse("rutas:document", args=[document.pk]))
        self.assertEqual(len(document.snapshot["labels"]), 2)
        self.assertEqual(document.snapshot["week"], "31A")
        self.assertEqual(len(PdfReader(io.BytesIO(document.pdf.open("rb").read())).pages), 1)


class ReconciliationTests(TestCase):
    def test_different_source_operations_remain_separate_and_reviewed(self):
        client = Client.objects.create(name="DAIKIN")
        one = SourceBook.objects.create(filename="uno.xlsm", sha256="a" * 64, classification="Headers")
        two = SourceBook.objects.create(filename="dos.xlsm", sha256="b" * 64, classification="SLP Headers")
        for source in (one, two):
            ImportRun.objects.create(source=source, filename=source.filename, sha256=source.sha256, dry_run=False)
        first = Route.objects.create(client=client, code="A-1", source=one, source_sheet="Ruta", source_row=2, status=Route.ACTIVE)
        second = Route.objects.create(client=client, code="A-1", source=two, source_sheet="Ruta", source_row=4, status=Route.ACTIVE)
        Operation.objects.create(route=first, position=1, source_sequence="1", name="CORTE", tooling="H-A")
        Operation.objects.create(route=second, position=1, source_sequence="1", name="CORTE", tooling="H-B")
        Part.objects.create(client=client, code="FP-1", part_type="FP", source=one, source_row=8)
        stats = reconcile_all()
        self.assertEqual(stats, {"crossbook_conflicts": 2, "missing_fp_routes": 1})
        self.assertEqual(reconcile_all(), {"crossbook_conflicts": 0, "missing_fp_routes": 0})
        self.assertEqual(Route.objects.filter(status=Route.REVIEW).count(), 2)
        self.assertEqual(Route.objects.count(), 2)
        reviewer = get_user_model().objects.create_superuser("revisor", "r@example.test", "example-long-password")
        self.client.force_login(reviewer)
        issue = ImportIssue.objects.get(run__source=one, sheet="Ruta", kind="diferencia_entre_libros")
        self.assertEqual(self.client.get(reverse("rutas:issues")).status_code, 200)
        self.client.post(reverse("rutas:issue_resolve", args=[issue.pk]), {"resolution_note": "Se compararon ambas variantes y se conserva Headers."})
        issue.refresh_from_db()
        self.assertTrue(issue.resolved)
        self.assertEqual(issue.resolved_by, reviewer)
        self.assertEqual(Route.objects.get(pk=first.pk).status, Route.REVIEW)
